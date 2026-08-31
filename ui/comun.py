"""
ui/comun.py — Utilidades y fachada de la UI

CSS, paleta, helpers de formato/descarga/controles y re-exportación de las
funciones de core que usan las páginas. Cada página hace `from ui.comun import *`.
"""
import io
import sys
from pathlib import Path
from datetime import date

import streamlit as st
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go

ROOT = Path(__file__).parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from setup_db import get_db_path, get_connection, modo_solo_lectura
from core.comun import (
    listar_categorias, listar_fechas_valorizacion,
    catalogo_articulos, filtrar_catalogo,
    tipos_movimiento, movimientos_detalle,
    MODOS_VALORIZACION, MODOS_VALORIZACION_MERMA,
)
from core.merma.analisis import (
    calcular_merma, merma_por_sucursal,
    calcular_merma_dinamica, merma_por_sucursal_dinamica,
    fechas_corte_comprobante,
)
from core.merma.avance import (
    avance_matriz, avance_conteo, NIVELES_AVANCE, METODOS_DENOMINADOR,
)
from core.merma.control import evolucion_mensual, movimientos_outliers, ajustes_por
from core.merma.ultimos_movimientos import ultimo_movimiento_por_tipo
from core.merma.transformaciones import detalle_transformaciones, pares_transformacion
from core.inventario.salud import analizar_stock, resumen_salud, ESTADOS
from core.inventario.series import (
    serie_stock, resumen_quiebres, serie_diaria, series_diarias,
    marcadores_movimiento, marcadores_sucursal, quiebres_por_dia,
    stock_por_sucursal, stock_por_categoria, TIPOS_MARCA, TIPOS_MARCA_SUC,
    resumen_movimientos, etiqueta_periodo,
    GRANULARIDADES, DIMENSIONES_RESUMEN, CLASIFICACIONES,
)
from core.inventario.interanual import (
    comparar_anios, serie_mensual_anios, serie_mensual_sucursales,
    ventas_vs_remitos, anios_disponibles, cobertura_meses,
    METRICAS as METRICAS_INTERANUAL, DIMENSIONES, MESES,
)
from core.inventario.politicas import (
    calcular_politicas, resumen_politicas, comparar_clase_abc, ESTADOS_POLITICA,
)
from core.inventario.reposicion import perfil_reposicion, resumen_reposicion, pedido_por_dias
from core.inventario.pedido_general import perfil_pedido_general, tabla_pedido_general
from core.distribucion.transferencias import sugerir_transferencias
from core.distribucion.diferencias import (
    diferencias_camion, resumen_diferencias, evolucion_diferencias,
    total_transferido, resumen_transferido, agregar_contraste_transferido,
    DEPOSITOS_DISTRIBUCION, FALTANTE, SOBRANTE, ENTRE_DEPOSITOS,
)
from core.comercial.margenes import analizar_margenes, evolucion_costos
from core.comercial.comprobantes import (
    resumen_comprobantes, evolucion_comprobantes, evolucion_comprobantes_sucursales,
    evolucion_comprobantes_multi, perfil_sku_en_comprobante, afinidad_skus,
    tipos_disponibles, METRICAS_COMPROBANTE,
)
from core.forecasting.demanda import (
    forecast_ventas, serie_mensual_real, estimar_demanda_stock_sku, NIVELES, METRICAS,
    MIN_MESES_KNN,
)
from core.reporte import (
    generar_reporte, generar_reporte_control, generar_reporte_salud,
    generar_reporte_politicas, generar_reporte_forecast, generar_reporte_transferencias,
    generar_reporte_diferencias, generar_reporte_interanual,
    generar_reporte_cruce, generar_reporte_series, generar_reporte_reposicion,
    generar_reporte_comprobantes,
    generar_reporte_avance_general, generar_reporte_avance_conteo,
    generar_reporte_pedido_sucursal, generar_reporte_transformaciones,
    cargar_branding, guardar_branding,
)
from core.auth import (
    listar_usuarios, crear_usuario, cambiar_password,
    set_acceso_proyecto, eliminar_usuario,
)

# Sentinel para "todas las sucursales" en los selectbox de la UI
TODAS = "__todas__"

# Paleta para categorías en los gráficos
PALETA = ["#ff6b6b", "#ffa94d", "#748ffc", "#64ffda", "#f783ac", "#a9e34b", "#ffd43b"]

# Layout de los gráficos, adaptado al tema activo de Streamlit (oscuro/claro).
# Se llaman en cada render (no constantes a nivel de módulo): ui/comun.py se
# importa una sola vez por proceso, así que un valor fijo quedaría congelado
# en el tema del primer render y no reaccionaría al toggle del usuario.
def layout_grafico() -> dict:
    """Layout base de los gráficos Plotly, adaptado al tema activo."""
    claro = st.context.theme.type == "light"
    return dict(
        plot_bgcolor="rgba(0,0,0,0)", paper_bgcolor="rgba(0,0,0,0)",
        font=dict(color="#1F2328" if claro else "#ccd6f6", family="IBM Plex Mono"),
        legend=dict(orientation="h", yanchor="bottom", y=1.02),
    )


def color_grilla() -> str:
    """Color de gridcolor para ejes, adaptado al tema activo. #d9dee6 es el
    mismo gris claro que usan los reportes imprimibles (core/graficos.py)."""
    return "#d9dee6" if st.context.theme.type == "light" else "#1e2130"

# CSS (se aplica desde app.py con st.markdown(CSS, unsafe_allow_html=True))
CSS = """
<style>
    /* Fuente principal */
    @import url('https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:wght@400;600&family=IBM+Plex+Sans:wght@300;400;600&display=swap');

    html, body, [class*="css"] {
        font-family: 'IBM Plex Sans', sans-serif;
    }

    /* Sidebar */
    [data-testid="stSidebar"] {
        background-color: #0f1117;
        border-right: 1px solid #1e2130;
    }
    [data-testid="stSidebar"] * {
        color: #e0e0e0 !important;
    }

    /* Métricas */
    [data-testid="stMetric"] {
        background: #1e2130;
        border: 1px solid #2d3148;
        border-radius: 8px;
        padding: 16px;
    }
    [data-testid="stMetricLabel"] {
        font-family: 'IBM Plex Mono', monospace;
        font-size: 0.7rem;
        letter-spacing: 0.1em;
        text-transform: uppercase;
        color: #8892b0 !important;
    }
    [data-testid="stMetricValue"] {
        font-family: 'IBM Plex Mono', monospace;
        font-size: 1.6rem;
        color: #ccd6f6 !important;
    }

    /* Encabezados */
    h1, h2, h3 {
        font-family: 'IBM Plex Mono', monospace;
        letter-spacing: -0.02em;
    }

    /* Tablas */
    [data-testid="stDataFrame"] {
        border: 1px solid #1e2130;
        border-radius: 8px;
    }

    /* Botones */
    .stButton > button {
        background: #1e2130;
        border: 1px solid #2d3148;
        color: #ccd6f6;
        font-family: 'IBM Plex Mono', monospace;
        font-size: 0.8rem;
        letter-spacing: 0.05em;
        border-radius: 4px;
        transition: all 0.2s;
    }
    .stButton > button:hover {
        background: #2d3148;
        border-color: #64ffda;
        color: #64ffda;
    }

    /* Selectbox */
    .stSelectbox label {
        font-family: 'IBM Plex Mono', monospace;
        font-size: 0.75rem;
        letter-spacing: 0.08em;
        text-transform: uppercase;
        color: #8892b0;
    }

    /* Tag de merma alta */
    .merma-alta { color: #ff6b6b; font-weight: 600; }
    .merma-media { color: #ffa94d; }
    .merma-baja { color: #64ffda; }
</style>
"""


def listar_proyectos() -> list[str]:
    projects_dir = ROOT / "projects"
    if not projects_dir.exists():
        return []
    return sorted([
        p.name for p in projects_dir.iterdir()
        if p.is_dir() and (p / "data.duckdb").exists()
    ])


def listar_sucursales(proyecto: str) -> pd.DataFrame:
    """Sucursales con movimientos, con nombre y abreviatura si existen."""
    conn = get_connection(proyecto)
    df = conn.execute("""
        SELECT DISTINCT m.codigodepo, d.nombre, d.abreviacion
        FROM movimientos m
        LEFT JOIN depositos d ON m.codigodepo = d.codigodepo
        WHERE m.codigodepo IS NOT NULL
        ORDER BY m.codigodepo
    """).df()
    conn.close()
    return df


def etiquetas_sucursal(df_suc: pd.DataFrame, tope: int = 28) -> tuple[dict, dict]:
    """
    Dos rótulos por sucursal a partir de `listar_sucursales`:

        largo → "008 — SAN JUAN"   para selectores y tablas (el código sirve
                                    para buscar y para cruzar con el ERP)
        corto → "San Juan"          para los gráficos, donde el código solo no
                                    dice nada

    El corto es el **nombre** cuando entra en `tope` caracteres. Si no entra,
    la abreviatura —pero sólo si es única, porque varias comparten "ZZZ"—; si
    tampoco, el nombre recortado. Cualquier empate que quede se resuelve
    agregando el código, así dos barras nunca comparten rótulo.
    """
    filas = []
    for _, r in df_suc.iterrows():
        nombre = str(r["nombre"]).strip() if pd.notna(r.get("nombre")) else ""
        abrev = (str(r["abreviacion"]).strip()
                 if pd.notna(r.get("abreviacion")) else "")
        filas.append((r["codigodepo"], " ".join(nombre.split()), abrev))

    unicas = {a for _, _, a in filas
              if a and sum(1 for _, _, x in filas if x == a) == 1}

    largo, base = {}, {}
    for c, nombre, abrev in filas:
        largo[c] = f"{c} — {nombre}".strip(" —") or str(c)
        if not nombre:
            base[c] = abrev or str(c)
        elif len(nombre) <= tope:
            base[c] = nombre.title()
        elif abrev in unicas:
            base[c] = abrev
        else:
            base[c] = nombre.title()[:tope].rstrip(" .,-") + "…"

    repetidos = {v for v in base.values()
                 if sum(1 for x in base.values() if x == v) > 1}
    return largo, {c: (f"{v} ({c})" if v in repetidos else v)
                   for c, v in base.items()}


def abreviaciones_sucursal(df_suc: pd.DataFrame) -> dict:
    """
    Mapa codigodepo -> abreviatura real de `depositos.abreviacion` (ej.
    "SJN"), a diferencia de `etiquetas_sucursal` que solo usa la
    abreviatura como respaldo cuando el nombre no entra en el tope de
    caracteres. Si falta la abreviatura, usa el codigodepo; si dos
    sucursales del propio `df_suc` comparten abreviatura, desambigua
    agregando el código (mismo patrón que `etiquetas_sucursal`).
    """
    base = {}
    for _, r in df_suc.iterrows():
        c = r["codigodepo"]
        abrev = (str(r["abreviacion"]).strip()
                 if pd.notna(r.get("abreviacion")) else "")
        base[c] = abrev or str(c)

    repetidos = {v for v in base.values()
                 if sum(1 for x in base.values() if x == v) > 1}
    return {c: (f"{v} ({c})" if v in repetidos else v)
            for c, v in base.items()}


def etiqueta_seleccion_sucursales(suc_sel: list[str], suc_largo: dict) -> str:
    """Texto descriptivo de una selección de sucursales (multiselect vacío =
    todas), para mostrar en metadatos y reportes."""
    if not suc_sel:
        return "Todas las sucursales"
    if len(suc_sel) == 1:
        return suc_largo.get(suc_sel[0], suc_sel[0])
    return f"{len(suc_sel)} sucursales: " + ", ".join(suc_sel)


def rango_disponible(proyecto: str) -> tuple[date | None, date | None]:
    """Rango de fechas con movimientos."""
    conn = get_connection(proyecto)
    rm = conn.execute("SELECT MIN(fecha), MAX(fecha) FROM movimientos").fetchone()
    conn.close()
    return rm[0], rm[1]


def categorias_merma(proyecto: str) -> list[str]:
    """Categorías marcadas es_merma, en orden de catálogo."""
    conn = get_connection(proyecto)
    rows = conn.execute("""
        SELECT categoria FROM tipos_categoria
        WHERE es_merma GROUP BY categoria ORDER BY MIN(orden), categoria
    """).fetchall()
    conn.close()
    return [r[0] for r in rows]


def listar_periodos_ingesta(proyecto: str) -> pd.DataFrame:
    conn = get_connection(proyecto)
    df = conn.execute("""
        SELECT tabla, periodo, archivo_origen, estado,
               registros_cargados, fecha_ingesta
        FROM periodos_ingesta
        ORDER BY fecha_ingesta DESC
    """).df()
    conn.close()
    return df


@st.cache_data(show_spinner=False)
def catalogo_cacheado(proyecto: str) -> pd.DataFrame:
    """Catálogo de SKUs con jerarquía resuelta (para armar los filtros)."""
    return catalogo_articulos(proyecto)


@st.cache_data(show_spinner=False)
def forecast_ventas_cacheado(proyecto: str, codigodepo: str | None, nivel: str,
                             horizonte: int, backtest: int, metrica: str) -> pd.DataFrame:
    """Cachea forecast_ventas — la competencia de modelos (Holt-Winters/SARIMA
    además del estacional robusto, por grupo) es bastante más cara que el
    modelo único de antes, sobre todo en niveles con muchos grupos (marca,
    rubro); evita recalcular al repetir los mismos parámetros."""
    return forecast_ventas(proyecto, codigodepo=codigodepo, nivel=nivel,
                           horizonte=horizonte, backtest=backtest, metrica=metrica)


def filtros_catalogo(cat: pd.DataFrame, key: str,
                     con_texto: bool = True,
                     con_activo: bool = False,
                     con_gsr: bool = True) -> tuple[pd.DataFrame, bool]:
    """
    Multiselects de Gran Super Rubro / Rubro / Marca (+ búsqueda por texto),
    encadenados: cada uno ofrece solo lo que sobrevive a los anteriores.

    `con_activo=True` agrega un checkbox "Ocultar SKUs inactivos", sin marcar
    por defecto — un SKU inactivo puede haber tenido ventas hasta hace poco,
    así que no se oculta solo.

    `con_gsr=False` oculta el filtro de Gran Super Rubro (páginas donde no se
    quiere exponer esa dimensión a la instancia web).

    Returns:
        (catálogo filtrado, hay_algun_filtro_activo)
    """
    def opciones(df, col):
        return sorted(df[col].dropna().astype(str).unique().tolist())

    cols = iter(st.columns(2 + con_gsr + con_texto + con_activo))

    gsr = []
    if con_gsr:
        with next(cols):
            gsr = st.multiselect("Gran Super Rubro", opciones(cat, "gran_super_rubro"),
                                 key=f"{key}_gsr")
    parcial = filtrar_catalogo(cat, gran_super_rubros=gsr)
    with next(cols):
        rubros = st.multiselect("Rubro", opciones(parcial, "rubro"), key=f"{key}_rub")
    parcial = filtrar_catalogo(parcial, rubros=rubros)
    with next(cols):
        marcas = st.multiselect("Marca", opciones(parcial, "marca"), key=f"{key}_mar")
    parcial = filtrar_catalogo(parcial, marcas=marcas)

    texto = ""
    if con_texto:
        with next(cols):
            texto = st.text_input("Buscar código / descripción", key=f"{key}_txt")
        parcial = filtrar_catalogo(parcial, texto=texto)

    ocultar_inactivos = False
    if con_activo:
        with next(cols):
            ocultar_inactivos = st.checkbox(
                "Ocultar SKUs inactivos", value=False, key=f"{key}_activo",
                help="Un SKU inactivo puede haber tenido ventas hasta hace poco.",
            )
        parcial = filtrar_catalogo(parcial, solo_activos=ocultar_inactivos)

    return (parcial,
            bool(gsr or rubros or marcas or (texto or "").strip() or ocultar_inactivos))


_ETIQUETAS_DIMENSION = {
    "gran_super_rubro": "Gran Super Rubro", "rubro": "Rubro", "marca": "Marca",
}


def filtros_resultado(
    df: pd.DataFrame, key: str,
    dimensiones: tuple[str, ...] = ("gran_super_rubro", "rubro", "marca"),
    con_texto: bool = True,
) -> tuple[pd.DataFrame, bool]:
    """
    Igual que filtros_catalogo(), pero sobre un DataFrame *ya calculado*
    (el resultado de un análisis) en vez del catálogo crudo — para las
    páginas que filtran su propio resultado sin recalcular. `dimensiones`
    controla qué columnas de jerarquía ofrecer (las que no estén en `df` se
    saltean), encadenadas en ese orden.

    Returns:
        (df filtrado, hay_algún_filtro_activo)
    """
    dims = [d for d in dimensiones if d in df.columns]
    cols = iter(st.columns(len(dims) + con_texto)) if (dims or con_texto) else iter([])

    parcial = df
    seleccion = {}
    for dim in dims:
        with next(cols):
            opciones = sorted(parcial[dim].dropna().astype(str).unique().tolist())
            sel = st.multiselect(_ETIQUETAS_DIMENSION.get(dim, dim), opciones,
                                 key=f"{key}_{dim}")
        seleccion[dim] = sel
        if sel:
            parcial = parcial[parcial[dim].astype(str).isin(sel)]

    texto = ""
    if con_texto:
        with next(cols):
            texto = st.text_input("Buscar código / descripción", key=f"{key}_txt")
        if texto and texto.strip():
            t = texto.strip().lower()
            parcial = parcial[
                parcial["codigo"].astype(str).str.lower().str.contains(t, na=False)
                | parcial["descripcion"].astype(str).str.lower().str.contains(t, na=False)
            ]

    return parcial, bool(any(seleccion.values()) or (texto or "").strip())


def formatear_pesos(valor) -> str:
    if pd.isna(valor) or valor is None:
        return "—"
    return f"$ {valor:,.0f}"


def formatear_pct(valor) -> str:
    if pd.isna(valor) or valor is None:
        return "—"
    return f"{valor:.2f}%"


def formatear_unidades(valor) -> str:
    if pd.isna(valor) or valor is None:
        return "—"
    return f"{valor:,.0f}"


def merma_por_categoria(df: pd.DataFrame, cat: str) -> pd.Series:
    """
    Aporte de una categoría a la merma: un faltante aporta positivo, un
    ajuste positivo resta. Las columnas de categorías es_merma ya vienen
    con este signo aplicado desde core.merma (ver _ensamblar), así que
    suman directo a la merma total.
    """
    if cat not in df.columns:
        return pd.Series(0.0, index=df.index)
    return df[cat]


def boton_reporte(html: str, filename: str, key: str) -> None:
    """Botón de descarga del reporte HTML imprimible + aviso de Ctrl+P."""
    st.download_button(
        "🖨 Descargar reporte (HTML imprimible)", html.encode("utf-8"),
        filename, "text/html", key=key,
    )
    st.caption("Abrilo en el navegador e imprimí con Ctrl+P (o guardá como PDF).")


def a_excel(df: pd.DataFrame) -> bytes:
    """Serializa un DataFrame a .xlsx en memoria."""
    buf = io.BytesIO()
    df.to_excel(buf, index=False, engine="openpyxl")
    return buf.getvalue()


_FORMATO_MONEDA = '"$" #,##0'
_FORMATO_1_DECIMAL = "0.0"


def a_excel_pedido_general(tabla: pd.DataFrame) -> bytes:
    """
    Serializa la tabla ancha de `tabla_pedido_general` (ver
    `core.inventario.pedido_general`) a .xlsx con encabezados agrupados y
    formato condicional — versión de prueba para "Pedido general de
    distribución 2", no la usa la página en producción.

    Columnas esperadas (mismo orden que arma `tabla_pedido_general`):
    codigo, descripcion, marca, rubro, uxb, costo, lista_1, vta_prom_*,
    stock_cd_<CD>, stock_*, pedido_* (uno por sucursal), pedido_consolidado.

    Fila 1: celdas combinadas "DATOS" (7 columnas de producto), "VENTA
    PROMEDIO", "STOCK" (incluye el stock del CD) y "PEDIDO RECOMENDADO"
    (solo las columnas por sucursal, no el consolidado).
    Fila 2: encabezados — costo/lista_1/uxb/pedido_consolidado con su
    nombre tal cual, y las columnas de venta promedio/stock/pedido con
    solo la abreviatura de sucursal (sin el prefijo vta_prom_/stock_/
    pedido_; el stock del CD queda como "CD"). Datos desde la fila 3.

    Formato numérico:
    - costo, lista_1: moneda (`_FORMATO_MONEDA`).
    - vta_prom_*, pedido_* (por sucursal) y pedido_consolidado: 1 decimal
      (`_FORMATO_1_DECIMAL`).

    Formato condicional (relleno negro, letra blanca):
    - Columnas de stock (stock_cd_* y stock_*): valor <= 0.
    - Columnas de pedido por sucursal (pedido_*, no pedido_consolidado): valor == -1.

    Bordes: línea horizontal fina común entre todas las filas de toda la
    tabla (desde la fila de títulos combinados hasta la última fila de
    datos), más un borde exterior grueso rodeando cada uno de los 4 grupos
    (DATOS, VENTA PROMEDIO, STOCK, PEDIDO RECOMENDADO) desde su celda
    combinada hacia abajo — pedido_consolidado no tiene grupo, así que
    solo lleva la línea fina común.
    """
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    cols_datos = ["codigo", "descripcion", "marca", "rubro", "uxb", "costo", "lista_1"]
    cols_datos = [c for c in cols_datos if c in tabla.columns]
    cols_vta = [c for c in tabla.columns if c.startswith("vta_prom_")]
    cols_stock = [c for c in tabla.columns if c.startswith("stock_")]
    cols_pedido = [c for c in tabla.columns
                   if c.startswith("pedido_") and c != "pedido_consolidado"]
    cols_moneda = [c for c in ("costo", "lista_1") if c in tabla.columns]
    cols_1_decimal = [*cols_vta, *cols_pedido,
                       *(["pedido_consolidado"] if "pedido_consolidado" in tabla.columns else [])]

    def _abrev(nombre, prefijo):
        return nombre[len(prefijo):]

    encabezados = {}
    for c in cols_vta:
        encabezados[c] = _abrev(c, "vta_prom_")
    for c in cols_stock:
        encabezados[c] = "CD" if c.startswith("stock_cd_") else _abrev(c, "stock_")
    for c in cols_pedido:
        encabezados[c] = _abrev(c, "pedido_")

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        tabla.to_excel(writer, index=False, startrow=1, sheet_name="Pedido general")
        ws = writer.sheets["Pedido general"]

        def col_idx(nombre):
            return tabla.columns.get_loc(nombre) + 1  # 1-indexado para openpyxl

        def combinar(cols, texto):
            if not cols:
                return
            c1, c2 = col_idx(cols[0]), col_idx(cols[-1])
            ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)
            celda = ws.cell(row=1, column=c1, value=texto)
            celda.alignment = Alignment(horizontal="center")
            celda.font = Font(bold=True)

        combinar(cols_datos, "DATOS")
        combinar(cols_vta, "VENTA PROMEDIO")
        combinar(cols_stock, "STOCK")
        combinar(cols_pedido, "PEDIDO RECOMENDADO")

        for nombre, texto in encabezados.items():
            ws.cell(row=2, column=col_idx(nombre), value=texto)

        fill_negro = PatternFill(start_color="FF000000", end_color="FF000000", fill_type="solid")
        font_blanco = Font(color="FFFFFFFF")
        n_filas = len(tabla)
        if n_filas:
            fila_desde, fila_hasta = 3, 2 + n_filas

            def rango(nombre):
                letra = get_column_letter(col_idx(nombre))
                return f"{letra}{fila_desde}:{letra}{fila_hasta}"

            for c in cols_moneda:
                for fila in range(fila_desde, fila_hasta + 1):
                    ws.cell(row=fila, column=col_idx(c)).number_format = _FORMATO_MONEDA
            for c in cols_1_decimal:
                for fila in range(fila_desde, fila_hasta + 1):
                    ws.cell(row=fila, column=col_idx(c)).number_format = _FORMATO_1_DECIMAL

            for c in cols_stock:
                ws.conditional_formatting.add(
                    rango(c),
                    CellIsRule(operator="lessThanOrEqual", formula=["0"],
                               fill=fill_negro, font=font_blanco),
                )
            for c in cols_pedido:
                ws.conditional_formatting.add(
                    rango(c),
                    CellIsRule(operator="equal", formula=["-1"],
                               fill=fill_negro, font=font_blanco),
                )

            fila_titulo, ultima_fila, ultima_col = 1, fila_hasta, len(tabla.columns)
            fina, gruesa = Side(style="thin", color="FF000000"), Side(style="thick", color="FF000000")

            # Línea horizontal fina común entre todas las filas, en toda la tabla.
            for fila in range(fila_titulo, ultima_fila + 1):
                for col in range(1, ultima_col + 1):
                    ws.cell(row=fila, column=col).border = Border(top=fina, bottom=fina)

            # Borde exterior grueso por grupo, desde la celda combinada hacia abajo.
            def borde_grupo(cols):
                if not cols:
                    return
                c1, c2 = col_idx(cols[0]), col_idx(cols[-1])
                for fila in range(fila_titulo, ultima_fila + 1):
                    for col in range(c1, c2 + 1):
                        celda = ws.cell(row=fila, column=col)
                        b = celda.border
                        celda.border = Border(
                            top=gruesa if fila == fila_titulo else b.top,
                            bottom=gruesa if fila == ultima_fila else b.bottom,
                            left=gruesa if col == c1 else b.left,
                            right=gruesa if col == c2 else b.right,
                        )

            for grupo in (cols_datos, cols_vta, cols_stock, cols_pedido):
                borde_grupo(grupo)
    return buf.getvalue()


def botones_descarga(df: pd.DataFrame, nombre: str, key: str) -> None:
    """Botones Excel + CSV para descargar un listado."""
    c1, c2, _ = st.columns([1, 1, 4])
    c1.download_button(
        "⬇ Excel", a_excel(df), f"{nombre}.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key=f"{key}_xlsx",
    )
    c2.download_button(
        "⬇ CSV", df.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig"),
        f"{nombre}.csv", "text/csv", key=f"{key}_csv",
    )


@st.cache_data(show_spinner=False)
def comparativa_cacheada(proyecto: str, fecha_desde: str, fecha_hasta: str,
                         modo_valorizacion: str, fecha_valorizacion,
                         codigos: tuple | None,
                         codigodepo: tuple | None = None) -> pd.DataFrame:
    return merma_por_sucursal(
        proyecto, fecha_desde, fecha_hasta,
        modo_valorizacion=modo_valorizacion,
        fecha_valorizacion=fecha_valorizacion,
        codigos=list(codigos) if codigos is not None else None,
        codigodepo=list(codigodepo) if codigodepo is not None else None,
    )


@st.cache_data(show_spinner=False)
def comparativa_dinamica_cacheada(proyecto: str, tipo_comprobante: str,
                                  cantidad: int, fecha_hasta: str,
                                  modo_valorizacion: str, fecha_valorizacion,
                                  codigos: tuple | None,
                                  codigodepo: tuple | None = None) -> pd.DataFrame:
    return merma_por_sucursal_dinamica(
        proyecto, tipo_comprobante, cantidad, fecha_hasta,
        codigodepo=list(codigodepo) if codigodepo is not None else None,
        modo_valorizacion=modo_valorizacion,
        fecha_valorizacion=fecha_valorizacion,
        codigos=list(codigos) if codigos is not None else None,
    )


@st.cache_data(show_spinner=False)
def fecha_corte_moda_cacheada(proyecto: str, tipo: str, cantidad: int,
                              fecha_hasta: str, codigodepo: tuple | None,
                              codigos: tuple | None):
    """Moda de fecha_corte entre los pares SKU+sucursal `suficientes`, para
    el subconjunto de SKUs efectivamente visible (post filtros de la UI)."""
    if not codigos:
        return None
    cortes = fechas_corte_comprobante(
        proyecto, tipo, cantidad, fecha_hasta,
        codigodepo=list(codigodepo) if codigodepo is not None else None,
        codigos=list(codigos),
    )
    cortes_ok = cortes[cortes["suficientes"]]
    if cortes_ok.empty:
        return None
    moda = pd.to_datetime(cortes_ok["fecha_corte"]).dt.date.mode()
    return moda.iloc[0] if not moda.empty else None


@st.cache_data(show_spinner=False)
def avance_matriz_cacheada(proyecto: str, fecha_desde: str, fecha_hasta: str,
                           nivel: str, metodo_denominador: str,
                           codigos: tuple | None,
                           codigodepo: tuple | None = None) -> pd.DataFrame:
    return avance_matriz(
        proyecto, fecha_desde, fecha_hasta, nivel=nivel,
        metodo_denominador=metodo_denominador,
        codigos=list(codigos) if codigos is not None else None,
        codigodepo=list(codigodepo) if codigodepo is not None else None,
    )


@st.cache_data(show_spinner=False)
def avance_conteo_cacheada(proyecto: str, fecha_desde: str, fecha_hasta: str,
                           modo_valorizacion: str, fecha_valorizacion,
                           codigos: tuple | None,
                           codigodepo: tuple | None = None) -> pd.DataFrame:
    return avance_conteo(
        proyecto, fecha_desde, fecha_hasta,
        modo_valorizacion=modo_valorizacion,
        fecha_valorizacion=fecha_valorizacion,
        codigos=list(codigos) if codigos is not None else None,
        codigodepo=list(codigodepo) if codigodepo is not None else None,
    )


def render_comparativa(df_suc: pd.DataFrame, merma_cats: list[str], key: str) -> None:
    """Tabla + gráfico + descarga de la comparativa por sucursal."""
    if df_suc.empty:
        st.info("Sin datos para la comparativa.")
        return

    df_suc = df_suc.copy()
    df_suc["sucursal"] = df_suc["codigodepo"] + " — " + df_suc["nombre"].fillna("")

    # Tabla
    cols = ["codigodepo", "nombre", "skus", "skus_con_merma"]
    headers = ["Código", "Sucursal", "SKUs", "SKUs c/merma"]
    formato = {}
    for c in merma_cats:
        if c in df_suc.columns:
            cols.append(c); headers.append(f"{c} $"); formato[c] = formatear_pesos
            cu = f"{c} (u)"
            cols.append(cu); headers.append(f"{c} u."); formato[cu] = formatear_unidades
    for c, h, f in [
        ("merma_total_valorizada", "Merma Total $", formatear_pesos),
        ("merma_total_unidades",   "Merma Total u.", formatear_unidades),
        ("venta_neta",             "Venta $",        formatear_pesos),
        ("pct_merma_sobre_ventas", "% Merma",        formatear_pct),
    ]:
        cols.append(c); headers.append(h); formato[c] = f

    df_display = df_suc[cols].copy()
    for c, f in formato.items():
        df_display[c] = df_display[c].apply(f)
    df_display.columns = headers
    st.dataframe(df_display, width="stretch", hide_index=True)

    botones_descarga(df_suc[cols], "comparativa_sucursales", key)

    # Gráfico: aporte a merma por categoría (con signo) + % en eje secundario
    fig = go.Figure()
    for i, cat in enumerate(merma_cats):
        if cat in df_suc.columns:
            fig.add_bar(
                name=cat, x=df_suc["sucursal"],
                y=df_suc[cat],
                marker_color=PALETA[i % len(PALETA)],
            )
    fig.add_scatter(
        name="% Merma s/venta", x=df_suc["sucursal"],
        y=df_suc["pct_merma_sobre_ventas"],
        mode="lines+markers", line=dict(color="#64ffda", width=2),
        marker=dict(size=7), yaxis="y2",
    )
    fig.update_layout(
        barmode="stack", height=430,
        xaxis=dict(tickangle=-30, gridcolor=color_grilla()),
        yaxis=dict(title="Merma $", gridcolor=color_grilla()),
        yaxis2=dict(title="% Merma", overlaying="y", side="right",
                    gridcolor=color_grilla(), ticksuffix="%"),
        margin=dict(b=120),
        **layout_grafico(),
    )
    st.plotly_chart(fig, width="stretch")


ETIQUETAS_MODO_VALORIZACION = {
    "costo": "A costo",
    "lista_1": "A precio de lista",
    "mixto": "Mixto (ventas a precio de venta, resto a costo)",
}


def controles_periodo(proyecto: str, key_prefix: str,
                      incluir_modo: bool = True,
                      desde_default: str = "min",
                      modos: list[str] | None = None):
    """
    Renderiza fecha_desde, fecha_hasta, [modo] y fecha de valorización.

    Args:
        incluir_modo:  si False, no muestra el selector "costo / lista_1"
                       (para páginas que no valorizan en $, p.ej. series de
                       stock en unidades) y devuelve "costo" sin preguntarlo.
        desde_default: "min" arranca en el primer movimiento disponible
                       (todo el histórico); "anio" arranca el 1° de enero
                       del año del último movimiento (recorte al año en
                       curso, más liviano para páginas de reconstrucción
                       diaria).
        modos:         opciones del selector de valorización. Default
                       `MODOS_VALORIZACION` (costo/lista_1). Merma pasa
                       `MODOS_VALORIZACION_MERMA` para sumar "mixto".
    """
    modos = list(modos) if modos else list(MODOS_VALORIZACION)
    fmin, fmax = rango_disponible(proyecto)
    if fmin is None:
        st.info("No hay movimientos cargados todavía.")
        return None

    snapshots = listar_fechas_valorizacion(proyecto)
    desde_ini = fmin if desde_default == "min" else max(fmin, fmax.replace(month=1, day=1))

    cols = iter(st.columns(3 + incluir_modo))
    with next(cols):
        fecha_desde = st.date_input(
            "Desde", value=desde_ini, min_value=fmin, max_value=fmax,
            key=f"{key_prefix}_desde",
        )
    with next(cols):
        fecha_hasta = st.date_input(
            "Hasta", value=fmax, min_value=fmin, max_value=fmax,
            key=f"{key_prefix}_hasta",
        )
    modo_val = "costo"
    if incluir_modo:
        with next(cols):
            modo_val = st.selectbox(
                "Valorización", modos,
                format_func=lambda x: ETIQUETAS_MODO_VALORIZACION.get(x, x),
                key=f"{key_prefix}_modo",
            )
    with next(cols):
        if snapshots:
            fecha_val = st.selectbox(
                "Valorizar al (snapshot)", snapshots,
                key=f"{key_prefix}_fval",
                help="Fecha del stock con la que se valoriza. Default: el más reciente.",
            )
        else:
            fecha_val = None
            st.warning("Sin snapshots de stock — no se puede valorizar.")

    return {
        "fecha_desde": str(fecha_desde),
        "fecha_hasta": str(fecha_hasta),
        "modo_valorizacion": modo_val,
        "fecha_valorizacion": fecha_val,
    }
